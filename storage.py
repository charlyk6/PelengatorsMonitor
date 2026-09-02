from pathlib import Path
from threading import Lock

from openpyxl import Workbook, load_workbook

EXCEL_PATH = Path(__file__).resolve().parent / "devices.xlsx"
HEADERS = ("Номер", "Описание")

_lock = Lock()


def _load_workbook():
    if not EXCEL_PATH.exists():
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Устройства"
        sheet.append(list(HEADERS))
        workbook.save(EXCEL_PATH)
        return workbook
    return load_workbook(EXCEL_PATH)


def list_devices() -> list[dict[str, str]]:
    with _lock:
        workbook = _load_workbook()
        sheet = workbook.active
        devices = []
        for number, description in sheet.iter_rows(min_row=2, values_only=True):
            if number is None:
                continue
            devices.append(
                {
                    "number": str(number).strip(),
                    "description": "" if description is None else str(description),
                }
            )
        return devices


def get_device(number: str) -> dict[str, str] | None:
    number = number.strip()
    for device in list_devices():
        if device["number"] == number:
            return device
    return None


def add_device(number: str) -> bool:
    number = number.strip()
    with _lock:
        workbook = _load_workbook()
        sheet = workbook.active
        for cell_number, _ in sheet.iter_rows(min_row=2, values_only=True):
            if cell_number is not None and str(cell_number).strip() == number:
                return False
        sheet.append([number, ""])
        workbook.save(EXCEL_PATH)
        return True


def update_description(number: str, description: str) -> bool:
    number = number.strip()
    with _lock:
        workbook = _load_workbook()
        sheet = workbook.active
        for row in sheet.iter_rows(min_row=2):
            cell_number = row[0].value
            if cell_number is not None and str(cell_number).strip() == number:
                row[1].value = description
                workbook.save(EXCEL_PATH)
                return True
        return False


def delete_device(number: str) -> bool:
    number = number.strip()
    with _lock:
        workbook = _load_workbook()
        sheet = workbook.active
        for row_index in range(2, sheet.max_row + 1):
            cell_number = sheet.cell(row=row_index, column=1).value
            if cell_number is not None and str(cell_number).strip() == number:
                sheet.delete_rows(row_index)
                workbook.save(EXCEL_PATH)
                return True
        return False
